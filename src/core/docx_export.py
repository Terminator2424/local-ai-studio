"""Exporta Markdown (con tablas, código y fórmulas LaTeX) a un documento de Word (.docx).

Las fórmulas se dibujan como imágenes (matplotlib mathtext, igual que en el visor); si alguna no es compatible se escribe en
Unicode legible. El texto procede de un PDF o de un modelo: se escribe como TEXTO en el documento (nunca como campos,
macros ni enlaces), y no se descarga nada.
"""
from __future__ import annotations

import io
import re
import unicodedata
from pathlib import Path

from .mdview import _MATH_SPAN, _formula_png, convert_arrays

_PLACE = re.compile(r"@@M(\d+)@@")
_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")          # caracteres que Word no admite en el XML


def _clean(s: str) -> str:
    return _CTRL.sub("", unicodedata.normalize("NFC", s))


def md_to_docx(md: str, path: str | Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt
    from markdown_it import MarkdownIt

    from ..pdf.textproc import latex_to_text

    md = _clean(convert_arrays(md))
    spans: list[tuple[str, bool]] = []

    def protect(m):
        body = next(g for g in m.groups() if g is not None).strip()
        spans.append((body, m.group(0).startswith(("$$", "\\["))))
        return f"@@M{len(spans) - 1}@@"

    tokens = MarkdownIt("commonmark", {"html": False}).enable("table").parse(_MATH_SPAN.sub(protect, md))
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    def add_math(par, idx: int):
        body, display = spans[idx]
        png = _formula_png(body, "#000000", 13.5 if display else 12.0)
        if png is None:
            run = par.add_run(re.sub(r"\s+", " ", latex_to_text(f"${body}$")).strip())
            run.italic = True
            return
        data, w, h = png
        height = Cm(max(0.45, min(2.2, h * 2.54 / 220 * 0.95)))
        par.add_run().add_picture(io.BytesIO(data), height=height)

    def add_text(par, text: str, bold=False, italic=False, code=False):
        for k, piece in enumerate(_PLACE.split(text)):
            if k % 2 == 1:
                add_math(par, int(piece))
            elif piece:
                run = par.add_run(piece)
                run.bold, run.italic = bold, italic
                if code:
                    run.font.name = "Consolas"
                    run.font.size = Pt(10)

    def render_inline(par, inline):
        bold = italic = False
        for ch in inline.children or []:
            t = ch.type
            if t == "text":
                add_text(par, ch.content, bold, italic)
            elif t == "strong_open":
                bold = True
            elif t == "strong_close":
                bold = False
            elif t == "em_open":
                italic = True
            elif t == "em_close":
                italic = False
            elif t == "code_inline":
                add_text(par, ch.content, code=True)
            elif t in ("softbreak", "hardbreak"):
                par.add_run().add_break()

    def block_math(par_text: str) -> int | None:
        m = _PLACE.fullmatch(par_text.strip())
        return int(m.group(1)) if m and spans[int(m.group(1))][1] else None

    i, list_stack = 0, []
    while i < len(tokens):
        tk = tokens[i]
        if tk.type == "heading_open":
            doc.add_heading("", level=min(int(tk.tag[1]), 4))
            render_inline(doc.paragraphs[-1], tokens[i + 1])
            i += 3
        elif tk.type == "paragraph_open":
            inline = tokens[i + 1]
            idx = block_math(inline.content)
            if idx is not None:                                   # fórmula en bloque: centrada en su propio párrafo
                par = doc.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_math(par, idx)
            else:
                style = "List Bullet" if list_stack and list_stack[-1] == "bullet" else \
                    "List Number" if list_stack and list_stack[-1] == "ordered" else None
                par = doc.add_paragraph(style=style)
                render_inline(par, inline)
            i += 3
        elif tk.type in ("bullet_list_open", "ordered_list_open"):
            list_stack.append("bullet" if tk.type.startswith("bullet") else "ordered")
            i += 1
        elif tk.type in ("bullet_list_close", "ordered_list_close"):
            list_stack.pop()
            i += 1
        elif tk.type in ("fence", "code_block"):
            par = doc.add_paragraph()
            run = par.add_run(tk.content.rstrip("\n"))
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            par.paragraph_format.left_indent = Cm(0.6)
            i += 1
        elif tk.type == "table_open":
            j = i
            rows: list[list] = []
            while tokens[j].type != "table_close":
                if tokens[j].type == "tr_open":
                    rows.append([])
                elif tokens[j].type in ("th_open", "td_open"):
                    rows[-1].append((tokens[j + 1], tokens[j].type == "th_open"))
                j += 1
            ncols = max(len(r) for r in rows) if rows else 0
            if ncols:
                table = doc.add_table(rows=len(rows), cols=ncols)
                table.style = "Table Grid"
                for r, row in enumerate(rows):
                    for c, (inline, head) in enumerate(row):
                        cell = table.cell(r, c)
                        cell.paragraphs[0].text = ""
                        render_inline(cell.paragraphs[0], inline)
                        if head:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
            i = j + 1
        elif tk.type == "hr":
            doc.add_paragraph("―" * 20)
            i += 1
        else:
            i += 1
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
